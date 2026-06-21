from __future__ import annotations

import io
import json
from pathlib import Path

from datetime import date, datetime

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload

from backend.config import Settings

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
SHEETS_SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
EXCEL_MIMES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}


def _credentials(settings: Settings, scopes: list[str]):
    return service_account.Credentials.from_service_account_file(
        str(settings.credentials_path),
        scopes=scopes,
    )


def _normalize_cell(cell):
    if cell is None:
        return ""
    if isinstance(cell, datetime):
        return cell.strftime("%Y-%m-%d")
    if isinstance(cell, date):
        return cell.isoformat()
    return cell


def _rows_from_worksheet(ws) -> list[list]:
    rows: list[list] = []
    for row in ws.iter_rows(values_only=True):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        rows.append([_normalize_cell(cell) for cell in row])
    return rows


class SheetsClient:
    """Native Google Sheets via Sheets API."""

    def __init__(self, settings: Settings):
        self.settings = settings
        self._service = None
        self._title = ""

    def _get_service(self):
        if self._service is None:
            creds = _credentials(self.settings, SHEETS_SCOPES)
            self._service = build("sheets", "v4", credentials=creds, cache_discovery=False)
        return self._service

    def get_spreadsheet_metadata(self) -> dict:
        service = self._get_service()
        result = (
            service.spreadsheets()
            .get(spreadsheetId=self.settings.spreadsheet_id, includeGridData=False)
            .execute()
        )
        self._title = result.get("properties", {}).get("title", "")
        sheets = []
        for sheet in result.get("sheets", []):
            props = sheet["properties"]
            grid = props.get("gridProperties", {})
            sheets.append(
                {
                    "title": props["title"],
                    "sheet_id": props["sheetId"],
                    "row_count": grid.get("rowCount", 0),
                    "col_count": grid.get("columnCount", 0),
                }
            )
        return {"title": self._title, "mime_type": GOOGLE_SHEET_MIME, "sheets": sheets}

    def get_sheet_values(self, sheet_title: str) -> list[list]:
        service = self._get_service()
        range_name = f"'{sheet_title}'"
        result = (
            service.spreadsheets()
            .values()
            .get(spreadsheetId=self.settings.spreadsheet_id, range=range_name)
            .execute()
        )
        return result.get("values", [])

    def fetch_all_tabs(self) -> dict[str, list[list]]:
        meta = self.get_spreadsheet_metadata()
        data: dict[str, list[list]] = {}
        for sheet in meta["sheets"]:
            data[sheet["title"]] = self.get_sheet_values(sheet["title"])
        return data


class ExcelDriveClient:
    """Excel-Datei auf Google Drive via Drive API + openpyxl."""

    def __init__(self, settings: Settings):
        self.settings = settings
        self._drive = None
        self._workbook = None
        self._meta: dict | None = None

    def _get_drive(self):
        if self._drive is None:
            creds = _credentials(self.settings, DRIVE_SCOPES)
            self._drive = build("drive", "v3", credentials=creds, cache_discovery=False)
        return self._drive

    def _load_workbook(self):
        if self._workbook is not None:
            return self._workbook

        drive = self._get_drive()
        file_id = self.settings.spreadsheet_id
        self._meta = (
            drive.files()
            .get(fileId=file_id, fields="name,mimeType")
            .execute()
        )

        request = drive.files().get_media(fileId=file_id)
        buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(buffer, request)
        done = False
        while not done:
            _, done = downloader.next_chunk()
        buffer.seek(0)
        self._workbook = openpyxl.load_workbook(buffer, data_only=True, read_only=True)
        return self._workbook

    def get_spreadsheet_metadata(self) -> dict:
        wb = self._load_workbook()
        sheets = []
        for idx, name in enumerate(wb.sheetnames):
            rows = _rows_from_worksheet(wb[name])
            sheets.append(
                {
                    "title": name,
                    "sheet_id": idx,
                    "row_count": len(rows),
                    "col_count": max((len(r) for r in rows), default=0),
                }
            )
        return {
            "title": self._meta["name"] if self._meta else "",
            "mime_type": self._meta.get("mimeType", "") if self._meta else "",
            "sheets": sheets,
        }

    def get_sheet_values(self, sheet_title: str) -> list[list]:
        wb = self._load_workbook()
        if sheet_title not in wb.sheetnames:
            return []
        return _rows_from_worksheet(wb[sheet_title])

    def fetch_all_tabs(self) -> dict[str, list[list]]:
        wb = self._load_workbook()
        return {name: _rows_from_worksheet(wb[name]) for name in wb.sheetnames}


class SampleSheetsClient:
    """Lädt Fixture-Daten für lokale Entwicklung ohne Google Credentials."""

    def __init__(self, fixture_path: Path | None = None):
        path = fixture_path or Path(__file__).resolve().parent.parent / "fixtures" / "sample_sheet.json"
        self._data = json.loads(path.read_text(encoding="utf-8"))

    def get_spreadsheet_metadata(self) -> dict:
        sheets = [
            {
                "title": name,
                "sheet_id": idx,
                "row_count": len(rows),
                "col_count": max((len(r) for r in rows), default=0),
            }
            for idx, (name, rows) in enumerate(self._data.items())
        ]
        return {"title": "Sample Budget Sheet", "mime_type": "sample", "sheets": sheets}

    def get_sheet_values(self, sheet_title: str) -> list[list]:
        return self._data.get(sheet_title, [])

    def fetch_all_tabs(self) -> dict[str, list[list]]:
        return dict(self._data)


def create_sheets_client(settings: Settings):
    if settings.use_sample_data:
        return SampleSheetsClient()
    if not settings.credentials_path.exists():
        raise FileNotFoundError(
            f"Google credentials not found at {settings.credentials_path}. "
            "Set USE_SAMPLE_DATA=true for local development."
        )

    creds = _credentials(settings, DRIVE_SCOPES)
    drive = build("drive", "v3", credentials=creds, cache_discovery=False)
    try:
        meta = (
            drive.files()
            .get(fileId=settings.spreadsheet_id, fields="name,mimeType")
            .execute()
        )
    except Exception as exc:
        raise RuntimeError(
            "Google Drive API Zugriff fehlgeschlagen. "
            "Drive API im GCP-Projekt aktivieren und die Excel-Datei "
            "mit dem Service Account teilen."
        ) from exc

    mime = meta.get("mimeType", "")
    if mime == GOOGLE_SHEET_MIME:
        return SheetsClient(settings)
    if mime in EXCEL_MIMES:
        return ExcelDriveClient(settings)

    raise ValueError(
        f"Unsupported file type: {mime}. Expected Google Sheet or Excel (.xlsx)."
    )
